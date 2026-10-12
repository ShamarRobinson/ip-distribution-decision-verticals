"""Stamp the Excel cost variance model with the current as-of date."""
import json
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font
ROOT = Path(__file__).resolve().parents[1]
AS = json.load(open(ROOT / "as_of.json"))
p = ROOT / "workbook" / "IP_Distribution_Cost_Variance_Model.xlsx"
wb = load_workbook(p)
wb["Summary"]["A4"] = f"Prices, fees and rates as of {AS['as_of_label']}. Refreshed quarterly; next refresh {AS['next_refresh_label']}."
wb["Summary"]["A4"].font = Font(name="Arial", italic=True, color="C00000")
src = wb["Sources"]
for row in src.iter_rows():
    for c in row:
        if isinstance(c.value, str) and c.value.startswith("Fees change often."):
            c.value = f"Fees change often. Figures are as of {AS['as_of_label']} and are refreshed quarterly. Not legal or financial advice."
wb.save(p)
print("stamped", AS["as_of_label"])
